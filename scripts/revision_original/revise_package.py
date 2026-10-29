"""Reproduce the descriptive reanalysis and editable Clinical Trials revision."""
from pathlib import Path
import csv
import difflib
import json
import re
import shutil
from collections import Counter

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OLD = ROOT / "manuscript/drug_safety_revision_manuscript.en.md"
DATA = HERE / "analysis"
DATA.mkdir(exist_ok=True)


def table_md(frame):
    def val(x):
        return "" if pd.isna(x) else str(x).replace("|", "/").replace("\n", " ")
    rows = [list(frame.columns)] + frame.values.tolist()
    return "\n".join("| " + " | ".join(val(x) for x in r) + " |" for r in rows)


def analyse():
    trials = pd.read_csv(ROOT / "tables/table1_included_trial_characteristics.csv")
    frames = []
    for is_structured in [True, False]:
        d = trials[(trials.fda_review_status == "structured values") == is_structured]
        frames.append({
            "Archived regulatory extraction": "Structured values" if is_structured else "Retrieved without structured core values",
            "Trials": len(d),
            "Registry values": sum(d.ctgov_status == "structured values"),
            "Randomized controlled": sum(d.design == "randomized, controlled"),
            "Other designs": sum(d.design != "randomized, controlled"),
            "Core outcomes median (range)": f"{d.core_concepts_reported_any_source_count.median():g} ({d.core_concepts_reported_any_source_count.min()}-{d.core_concepts_reported_any_source_count.max()})",
        })
    group = pd.DataFrame(frames)
    group.to_csv(DATA / "six_vs_seventeen.csv", index=False)
    pairs = pd.read_csv(ROOT / "tables/confirmed_pair_structured_audit_packet.csv")
    pairs["revision_use"] = "Descriptive count check; timing alignment not independently established"
    pairs.loc[pairs.final_current_analysis_tier == "sensitivity_candidate", "revision_use"] = "Contextual comparison; retain population or timing differences"
    pairs.loc[pairs.comparison_id == "FULLCOMP000137", "revision_use"] = "Reclassified; publication death explicitly outside safety reporting period"
    pairs.to_csv(DATA / "reappraised_28_record_inventory.csv", index=False)
    assert len(pairs) == 28
    assert sum(pairs.final_current_analysis_tier == "primary_candidate") == 21
    assert sum(pairs.comparison_id == "FULLCOMP000137") == 1
    assert len(trials) == 23
    assert group.Trials.tolist() == [6, 17]
    joint = pd.read_csv(ROOT / "tables/five_state_source_reporting_status.csv")
    use = joint[joint.five_state_status.isin(["reported_but_non_comparable", "comparable_and_concordant", "comparable_but_discordant"])]
    summary = use.groupby("source_pair").agg(**{"Screening groups": ("trial_id", "size"), "Trials contributing": ("trial_id", "nunique")}).reset_index()
    summary.to_csv(DATA / "legacy_screening_groups.csv", index=False)
    return trials, pairs, group, summary


def make_supplement(trials, pairs, group, summary):
    lines = ["# Supplementary information for CT-26-0272", "", "## Scope and revision audit", "",
        "This supplement retains all 23 trials and documents the original screening units. Original classifications describe the archived workflow, not independently validated interchangeability. All counts in this supplement were reproduced from the project tables during revision.", "",
        "## Supplementary Table S1 Included trials and archived source status", ""]
    t = trials[["short_trial_name", "nct_number", "phase", "design", "ctgov_status", "fda_review_status", "core_concepts_reported_any_source_count"]].copy()
    t.columns = ["Trial", "Registry identifier", "Phase", "Design", "Registry", "Regulatory extraction", "Core outcomes"]
    t["Trial"] = [f"{name} [{i}]" for i, name in enumerate(t.Trial, 1)]
    t["Registry"] = t.Registry.replace({"structured values": "Values", "not linked/unavailable": "No extracted values"})
    t["Regulatory extraction"] = t["Regulatory extraction"].replace({"structured values": "Values", "retrieved, no structured core values": "Retrieved; no core extraction"})
    lines += [table_md(t), "", "All trials had an extracted publication by design. Core outcomes means the number of seven prespecified outcomes recorded in at least one source. A missing extraction is not evidence that a public document contains no safety information. DREAMM-2 included randomized dose cohorts without a conventional control arm; its original cohort label is retained in the archive. References identify the primary publications.", "",
        "## Supplementary Table S2 What the original screening counts mean", "",
        "A screening group means one trial, one requested outcome and two sources. For example, DREAMM-2 serious adverse events in its publication and regulatory review form one group. Its two dose cohorts can generate two numerical records. A record is a pairing of two source rows for closer review. A group is not a trial, an independent observation, or a three-source match.", "",
        table_md(pd.DataFrame([
            ["483", "Design grid", "23 trials x 7 requested outcomes x 3 source pairs; not 483 observed comparisons"],
            ["53", "Groups with related extracted values", "Included mortality records initially screened against fatal adverse events; related labels did not establish a common endpoint"],
            ["17", "Groups contributing retained records", "Included both original primary and sensitivity records; the old label directly comparable was too broad"],
            ["269", "Automatically generated source-row pairings", "Included multiple arms, doses and repeated rows; an algorithm-dependent workload count"],
            ["28", "Retained source-row records", "21 originally called primary and 7 sensitivity; preserved here with revision annotations"],
            ["1", "Original primary record reclassified", "FULLCOMP000137 refers to a publication death explicitly outside the safety reporting period"],
        ], columns=["Count", "Unit", "Interpretation"])), "", table_md(summary), "",
        "Five trials had structured values in all three source types. The 17 groups were distributed across five trials, rather than representing complete three-source matches in each trial. They comprised 13 publication-regulatory groups, two publication-registry groups and two regulatory-registry groups. These historical counts include contextual comparisons and should not be used as a current direct-comparability rate.", "",
        "## Supplementary Table S3 Reappraisal of the retained source-row records", "",
        "The original tier is retained for traceability. Reclassification does not certify the remaining records as aligned in all dimensions. The revised manuscript uses source-specific values as descriptive illustrations and reports no pooled agreement estimate.", ""]
    names = trials.set_index("trial_id").short_trial_name.to_dict()
    for trial, frame in pairs.groupby("trial_id", sort=False):
        lines += [f"### {names[trial]}", ""]
        rows = []
        for _, r in frame.iterrows():
            v = lambda n, d, pc: f"{int(n)}/{int(d)} ({pc:g}%)" if pd.notna(n) else f"{pc:g}%"
            status="Count check; timing unverified"
            if r.final_current_analysis_tier=="sensitivity_candidate":status="Contextual comparison"
            if r.comparison_id=="FULLCOMP000137":status="Reclassified; outside reporting period"
            rows.append([r.comparison_id.replace("FULLCOMP", ""), r.ae_concept.replace("_", " "), r.source_pair.replace("FDA review", "regulatory review"), v(r.number_patients_1,r.denominator_1,r.percentage_1),v(r.number_patients_2,r.denominator_2,r.percentage_2),"Primary" if r.final_current_analysis_tier == "primary_candidate" else "Sensitivity",status])
        lines += [table_md(pd.DataFrame(rows,columns=["Record", "Outcome", "Sources in order", "Source 1", "Source 2", "Original tier", "Revision interpretation"])), ""]
    lines += ["## Supplementary Table S4 Worked examples and source locators", ""]
    for cid in ["FULLCOMP000003", "FULLCOMP000106", "FULLCOMP000119", "FULLCOMP000137", "FULLCOMP000150"]:
        r = pairs.loc[pairs.comparison_id == cid].iloc[0]
        lines += [f"### {cid} {names[r.trial_id]}", "", f"Source 1 term: {r.term_1}. Locator: {r.locator_1}. Population: {r.analysis_population_1}.", "", f"Source 2 term: {r.term_2}. Locator: {r.locator_2}. Population: {r.analysis_population_2}.", ""]
    lines += ["## Registry metadata checks", "",
        "The archived DREAMM-2 record (NCT03525678), adverseEventsModule.timeFrame, specifies collection from treatment start through a maximum of 186 weeks in the main study and approximately 325 weeks in the later phase. Its 2.5 mg/kg main-study group records 43/95 participants with serious adverse events. The publication and regulatory review record 38/95. The available metadata support a difference in reporting context; they do not isolate the contribution of follow-up duration to the five additional participants.", "",
        "For the same DREAMM-2 registry group, deaths were recorded as 70/97, whereas serious adverse events used a denominator of 95. The registry description explicitly includes all randomized participants for all-cause mortality and treated participants for adverse events. An identical group title therefore does not guarantee a common denominator across outcomes.", "",
        "The archived DESTINY-Breast01 registry record (NCT03248492) reports separate groups of 50, 48, 21, 130 and 4 participants. Its Part 1 low-dose group is not synonymous with the publication population of 184 participants. Its deaths field reports deaths from any cause, not deaths attributed to adverse events.", "",
        "The original outcome-measures screen produced 492 candidate rows and nine high-priority trial-outcome labels. These are search candidates, not nine confirmed additional aggregate outcomes. One candidate concerns patient-reported symptomatic adverse events and several concern laboratory abnormalities. No candidate was promoted to a verified aggregate endpoint during this revision.", "",
        "## Reproducibility and provenance", "",
        "The reanalysis uses tables/table1_included_trial_characteristics.csv, tables/confirmed_pair_structured_audit_packet.csv, tables/five_state_source_reporting_status.csv and the archived registry records in data/raw/clinicaltrials. The analysis/ directory contains the six-versus-seventeen summary, all 28 reappraised records and the historical screening-group counts. Simulated reviewer records were not used as evidence of independent adjudication. No new participant-level data, person-time denominators, independent duplicate adjudication or causal reconciliation were generated.", "",
        "The current revision scripts and outputs accompany this local supplement. The previously published repository links in the manuscript refer to the original project archive; the revision has not been uploaded to those services."]
    return "\n".join(lines) + "\n"


def prepare_references(text):
    original = OLD.read_text()
    refs = {int(n): s.strip() for n,s in re.findall(r"^(\d+)\. (.+)$", original, re.M)}
    refs[48] = "National Institutes of Health, Department of Health and Human Services. Clinical trials registration and results information submission. Final rule. Fed Regist 2016; 81: 64982-65157. https://www.govinfo.gov/content/pkg/FR-2016-09-21/html/2016-22129.htm."
    refs[49] = "US Food and Drug Administration. Premarketing risk assessment. Guidance for industry. March 2005. https://www.fda.gov/regulatory-information/search-fda-guidance-documents/premarketing-risk-assessment."
    refs[50] = "ClinicalTrials.gov. Results data element definitions. National Library of Medicine. https://clinicaltrials.gov/policy/results-definitions (accessed 24 October 2026)."
    refs[51] = "US Food and Drug Administration. Statistical considerations for premarketing risk assessment. Webinar, 16 May 2024. https://www.fda.gov/drugs/news-events-human-drugs/statistical-considerations-premarketing-risk-assessment-05162024."
    order = []
    def replace(m):
        ids = [int(x) for x in m.group(1).split(",")]
        for i in ids:
            if i not in order:
                order.append(i)
        return "[" + ",".join(str(i) for i in sorted(order.index(n)+1 for n in ids)) + "]"
    text = re.sub(r"\[([\d,]+)\]", replace, text)
    # Trial publications are cited in the accompanying supplement, after main text citations.
    for i in range(1,24):
        if i not in order:
            order.append(i)
    def format_ref(s):
        s = re.sub(r" PMID:.*?(?= DOI:|$)", "", s)
        s = re.sub(r" DOI: .+$", "", s) if "doi:" in s else s.replace(" DOI:", " doi:")
        s = re.sub(r" Epub [^.]+\.", "", s)
        # Use the first three authors followed by et al. in long journal entries.
        if ". " in s and not s.startswith(("US ", "National ", "ClinicalTrials")):
            authors, rest = s.split(". ", 1)
            parts = authors.split(", ")
            if len(parts) > 3:
                authors = ", ".join(parts[:3]) + ", et al"
            s = authors + ". " + rest
        return s
    reference_text = "\n## References\n\n" + "\n".join(f"{i+1}. {format_ref(refs[n])}" for i,n in enumerate(order)) + "\n\n"
    text = text.replace("## Table 1 ", reference_text + "## Table 1 ", 1)
    (DATA / "reference_number_map.json").write_text(json.dumps({str(n):i+1 for i,n in enumerate(order)},indent=2))
    return text, {i:order.index(i)+1 for i in range(1,24)}


def setup(doc, manuscript=False):
    sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.8)
    sec.left_margin=sec.right_margin=Inches(.85)
    for name in ["Normal","Title","Heading 1","Heading 2","Heading 3"]:
        st=doc.styles[name]; st.font.name="Times New Roman"; st.font.size=Pt(12)
        st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.line_spacing=2 if manuscript else 1.15
    doc.styles["Title"].font.size=Pt(16)
    doc.styles["Title"].font.bold=True
    for name in ["Heading 1","Heading 2","Heading 3"]:
        doc.styles[name].paragraph_format.keep_with_next=True
        doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles["Heading 2"].font.italic=True
    doc.styles["Heading 2"].font.bold=False
    for st in doc.styles:
        for border in st.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=foot.add_run(); fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); run._r.addnext(fld)
    if manuscript:
        ln=OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"),"1"); ln.set(qn("w:restart"),"continuous"); ln.set(qn("w:distance"),"240"); sec._sectPr.append(ln)


def normalise(s):
    return re.sub(r"\s+"," ",re.sub(r"\[[\d, -]+\]","",s)).strip()


def runs(p, s, red=False, refs=True):
    for part in re.split(r"(\[[\d, -]+\]|\*[^*]+\*)",s):
        if not part: continue
        if refs and re.fullmatch(r"\[[\d, -]+\]",part):
            r=p.add_run(part[1:-1]); r.font.superscript=True
        elif part.startswith("*") and part.endswith("*"):
            r=p.add_run(part[1:-1]); r.italic=True
        else: r=p.add_run(part)
        r.font.color.rgb=RGBColor.from_string("C00000" if red else "000000")


def add_table(doc, rows, red=False):
    t=doc.add_table(rows=1,cols=len(rows[0])); t.style="Table Grid"; t.autofit=False
    width=doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin
    weights=[1]*len(rows[0])
    if len(weights)==3: weights=[1.0,1.65,2.0]
    if len(weights)==4: weights=[1.05,1.2,1.2,1.8]
    if len(weights)==6: weights=[1.5,.5,.8,.8,.7,1.2]
    if len(weights)==7:
        weights=[1.4,1.15,.4,1.25,.9,1.45,.65] if "Phase" in rows[0] else [.65,1.45,1.6,1,1,.75,1.45]
    for j,w in enumerate(weights): t.columns[j].width=int(width*w/sum(weights))
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,cell in enumerate(cells):
            cell.width=int(width*weights[j]/sum(weights))
            p=cell.paragraphs[0]; p.paragraph_format.line_spacing=1.05; p.paragraph_format.space_after=Pt(4); p.paragraph_format.space_before=Pt(4)
            runs(p,row[j] if j<len(row) else "",red)
            for r in p.runs: r.font.size=Pt(9 if len(rows[0])>4 else 10); r.bold=i==0
            if i==0:
                sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),"EEEEEE"); cell._tc.get_or_add_tcPr().append(sh)
        trpr=t.rows[i]._tr.get_or_add_trPr()
        nosplit=OxmlElement("w:cantSplit"); trpr.append(nosplit)
        if i==0:
            rep=OxmlElement("w:tblHeader"); trpr.append(rep)
    doc.add_paragraph()
    borders=OxmlElement("w:tblBorders")
    for side in ["top","left","bottom","right","insideH","insideV"]:
        edge=OxmlElement("w:"+side);edge.set(qn("w:val"),"single");edge.set(qn("w:sz"),"4");edge.set(qn("w:color"),"D9D9D9");borders.append(edge)
    t._tbl.tblPr.append(borders)


def render_md(text,path,marked=False,manuscript=False,supplement=False):
    doc=Document(); setup(doc,manuscript)
    oldparas={normalise(p.text) for p in Document(HERE/"baseline/submitted_manuscript.docx").paragraphs if p.text.strip()}
    lines=text.splitlines(); i=0; in_abs=False
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                row=[x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?",x) for x in row): rows.append(row)
                i+=1
            add_table(doc,rows,marked); continue
        red=marked and normalise(re.sub(r"^#+ ","",line)) not in oldparas
        if line.startswith("# "):
            p=doc.add_paragraph(style="Title"); runs(p,line[2:],red)
        elif line.startswith("## "):
            title=line[3:]
            if title.startswith(("Table ","Figure ")):
                doc.add_page_break()
            elif supplement and title.startswith("Supplementary Table"):
                sec=doc.add_section(WD_SECTION.NEW_PAGE); sec.orientation=WD_ORIENT.LANDSCAPE
                sec.page_width=Inches(11);sec.page_height=Inches(8.5)
                sec.top_margin=sec.bottom_margin=Inches(.6)
            elif manuscript and title=="Abstract":
                doc.add_page_break(); in_abs=True
            elif manuscript and title=="Introduction":
                doc.add_page_break(); in_abs=False
            elif title=="Editor requirements":
                doc.add_page_break()
            p=doc.add_paragraph(style="Heading 1"); runs(p,title,red)
        elif line.startswith("### "):
            p=doc.add_paragraph(style="Heading 2"); runs(p,line[4:],red)
            if in_abs:
                for r in p.runs: r.italic=False;r.bold=True
        elif line.startswith("!["):
            image=HERE/re.search(r"\]\((.+)\)",line).group(1)
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(image),width=Inches(6.65))
        else:
            p=doc.add_paragraph(); runs(p,line,red)
        i+=1
    doc.core_properties.author="Da Lin; Yu Zhang"
    doc.core_properties.title=text.splitlines()[0].lstrip("# ")
    doc.save(path)


def figures(trials):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    assets=HERE/"figures";assets.mkdir(exist_ok=True)
    plt.rcParams.update({"font.family":"DejaVu Sans","font.size":10})
    fig,ax=plt.subplots(figsize=(10,6));ax.set_xlim(0,10);ax.set_ylim(0,6);ax.axis("off")
    labels=[("Define the safety question", "Trial, treatment group, outcome and observation period"),
        ("Identify what each source contributes", "Publication context | registry summaries | regulatory evaluation"),
        ("Check the underlying definitions", "Population, denominator, attribution, seriousness and timing"),
        ("Choose an appropriate interpretation", "Aligned count check | contextual comparison | additional information needed")]
    for i,(a,b) in enumerate(labels):
        y=4.75-i*1.35
        ax.add_patch(Rectangle((.25,y),9.5,.95,facecolor=["#E8EFF4","#F2F1E8","#E8F0EA","#EEEAF1"][i],edgecolor="#555555",linewidth=.7))
        ax.text(5,y+.61,a,ha="center",va="center",weight="bold",fontsize=12)
        ax.text(5,y+.25,b,ha="center",va="center",fontsize=10)
        if i<3: ax.annotate("",xy=(5,y-.35),xytext=(5,y-.05),arrowprops={"arrowstyle":"->","color":"#333333"})
    fig.tight_layout(pad=.5);fig.savefig(assets/"figure1_evidence_evaluation.png",dpi=220);plt.close(fig)
    import numpy as np
    a=np.array([[1,int(r.ctgov_status=="structured values"),int(r.fda_review_status=="structured values")] for _,r in trials.iterrows()])
    from matplotlib.colors import ListedColormap
    fig,ax=plt.subplots(figsize=(9,7))
    ax.imshow(a,cmap=ListedColormap(["#F0F0F0","#387568"]),vmin=0,vmax=1,aspect="auto")
    ax.set_xticks(range(3),["Publication\n23/23","ClinicalTrials.gov\n19/23","Regulatory review\n6/23"])
    ax.xaxis.tick_top();ax.tick_params(axis="both",length=0,pad=8)
    ax.set_yticks(range(23),trials.short_trial_name,fontsize=9)
    for i in range(23):
        for j in range(3):ax.text(j,i,"Yes" if a[i,j] else "No",ha="center",va="center",color="white" if a[i,j] else "#333333",fontsize=8)
    for s in ax.spines.values():s.set_visible(False)
    ax.set_xticks(np.arange(-.5,3,1),minor=True);ax.set_yticks(np.arange(-.5,23,1),minor=True);ax.grid(which="minor",color="white",linewidth=1)
    ax.tick_params(which="minor",length=0)
    fig.tight_layout();fig.savefig(assets/"figure2_source_inventory.png",dpi=220);plt.close(fig)


def main():
    baseline=HERE/"baseline";baseline.mkdir(exist_ok=True)
    originals={"submitted_manuscript.docx":ROOT/"manuscript/drug_safety_revision_manuscript.docx","submitted_manuscript.en.md":OLD,"CT-26-0272_Proof_hi.pdf":Path("/Users/zy/Downloads/CT-26-0272_Proof_hi.pdf"),"decision_letter.txt":Path("/Users/zy/.codex/attachments/adc9be7f-05a1-44ca-b2d7-6430187280ff/已粘贴的文本.txt")}
    for name,source in originals.items():
        if not (baseline/name).exists():shutil.copy2(source,baseline/name)
    trials,pairs,group,summary=analyse()
    figures(trials)
    raw=(HERE/"revised_manuscript.source.md").read_text()
    raw=raw.replace("{{GROUP_TABLE}}",table_md(group))
    text,trialrefs=prepare_references(raw)
    (HERE/"revised_manuscript.en.md").write_text(text)
    supp=make_supplement(trials,pairs,group,summary)
    def mapped(m):return "["+str(trialrefs[int(m.group(1))])+"]"
    supp=re.sub(r"\[(\d+)\]",mapped,supp)
    (HERE/"supplementary_information.en.md").write_text(supp)
    render_md(text,HERE/"CT-26-0272_revised_clean.docx",manuscript=True)
    render_md(text,HERE/"CT-26-0272_revised_marked.docx",marked=True,manuscript=True)
    render_md(supp,HERE/"CT-26-0272_supplementary_information.docx",supplement=True)
    response=(HERE/"response.source.md").read_text()
    render_md(response,HERE/"CT-26-0272_cover_letter_and_response.docx")
    clean=Document(HERE/"CT-26-0272_revised_clean.docx")
    marked=Document(HERE/"CT-26-0272_revised_marked.docx")
    assert clean._element.xpath("//w:t/text()") == marked._element.xpath("//w:t/text()")
    assert not clean._element.xpath('//w:color[@w:val="C00000"]')
    assert marked._element.xpath('//w:color[@w:val="C00000"]')
    abstract=text.split("## Abstract")[1].split("## Keywords")[0]
    body=text.split("\n## Introduction\n")[1].split("\n## Supplementary information\n")[0]
    wc=lambda s:len(re.findall(r"\b[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)*\b",re.sub(r"\[[\d, -]+\]","",s)))
    import platform
    intro=body.split("\n## Methods\n")[0]
    discussion=body.split("\n## Discussion\n")[1]
    metrics={"abstract_words_including_headings":wc(abstract),"main_words_including_headings":wc(body),"introduction_words":wc(intro),"discussion_words":wc(discussion),"discussion_to_introduction_ratio":round(wc(discussion)/wc(intro),2),"main_figures":2,"main_tables":3,"references":len(re.findall(r"^\d+\. ",text.split("## References")[1],re.M)),"clean_marked_text_identical":True,"python":platform.python_version(),"pandas":pd.__version__}
    (DATA/"package_checks.json").write_text(json.dumps(metrics,indent=2))
    print(json.dumps(metrics,indent=2))


if __name__=="__main__":main()
