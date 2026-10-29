"""Build a bounded R2 revision without modifying archived inputs."""
from pathlib import Path
import csv, json, re, difflib, importlib.util, copy, hashlib
from collections import Counter
from docx import Document
from docx.shared import RGBColor

HERE=Path(__file__).resolve().parent
PREV=HERE.parent/'CT-26-0272_revision'
ROOT=HERE.parents[1]
DATE='2026-10-25'
def load(n):return json.loads((HERE/'analysis'/n).read_text())
def save(n,x):(HERE/'analysis'/n).write_text(json.dumps(x,ensure_ascii=False,indent=2))
def table(rows):
    cols=list(rows[0]);clean=lambda x:str(x).replace('|','/').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(cols)+' |','| '+' | '.join(['---']*len(cols))+' |']+['| '+' | '.join(clean(r.get(c,'')) for c in cols)+' |' for r in rows])
def words(s):return len(re.findall(r"\b[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)*\b",re.sub(r'\[[\d, -]+\]','',s)))
def replace(text,old,new):
    assert old in text,old[:90]
    return text.replace(old,new,1)

def date_evidence():
    p=HERE/'source_date_followup/DREAMM8_historical_comparison.json'
    if not p.exists():return None
    d=json.loads(p.read_text())
    assert d['pre_cutoff_status']['lastUpdatePostDateStruct']['date']=='2025-03-18'
    assert d['pre_cutoff_equals_later_ae_module'] and d['archived_ae_semantically_equivalent']
    assert len(d['extracted_seed_checks'])==6 and all(r['matches_pre_cutoff'] for r in d['extracted_seed_checks'])
    return d

def augment_audit():
    rows=load('record_audit.json');trials=load('trial_matrix.json')
    refs=list(csv.DictReader((ROOT/'tables/primary_trial_reference_list.csv').open()))
    refby={x['short_trial_name']:x for x in refs}
    source_index={
        'DREAMM2_publication':('Main article safety text; supplement Table S1','Safety population 95 and 99; 21 June 2019 cutoff; dose modification, serious and fatal counts checked in main text'),
        'DREAMM2_FDA':('Table 27, PDF p. 148; Table 30, PDF p. 150','Table 27: 8/95, 10/99; 93/95, 99/99; 51/95, 61/99; 28/95, 41/99; 3/95, 7/99; 38/95, 47/99. Table 27 displays 47 (48), Table 30 displays 47 (47).'),
        'EV201_publication':('Table 3; Appendix Tables A3 and A4; safety-results narrative','N=125; treatment-related any 117 (94), treatment-related grade >=3 68 (54); all-cause any 125 (100), grade >=3 91 (73). One treatment-related death was outside the 30-day safety reporting period.'),
        'EV201_FDA':('Table 74 and footnotes, PDF pp. 184-185','Cohort 1 N=125: TEAE 125 (100); drug-related 117 (93.6); grade >=3 91 (72.8); drug-related grade >=3 68 (54.4); serious 58 (46.4); drug-related fatal 0. Footnote excludes the death >30 days after last dose.'),
        'IMMU13201_appendix':('Supplementary Table S1, PDF p. 18','108 participants with metastatic triple-negative breast cancer; adverse event as reason for discontinuation: 3 (2.8).'),
        'IMMU13201_FDA':('Discontinuations due to adverse effects, PDF pp. 87-88','mTNBC N=108; five (4.6%) participants discontinued due to TEAEs. Individual reviewer assessments do not explicitly reconcile this with the publication three.'),
        'DESTINY_appendix':('Supplementary Table S3, PDF p. 12 (printed p. 11)','N=184; any TEAE 183 (99.5), interruption 65 (35.3), serious 42 (22.8).'),
        'DESTINY_FDA':('Tables 35, 38 and 41; HER2-positive breast cancer 5.4 mg/kg pool','N=234 pool includes DS8201-A-J101 and DS8201-A-U201: serious 47 (20.1), interruption 78 (33.3), any TEAE 233 (99.6).'),
        'innovaTV301_appendix':('Supplementary Table S7, PDF p. 36 (printed p. 35); safety definitions','Tisotumab vedotin N=250: 82 (32.8) serious TEAEs; chemotherapy N=239: 94 (39.3). TEAEs through 30 days after last dose; publication cutoff 24 July 2023.'),
    }
    for r in rows:
        short=r['RECORD_ID'][-6:];name=r['TRIAL_NAME'];ref=refby[name]
        r['DOI']=ref['doi'];r['NCT_IDENTIFIER']=ref['nct_number']
        r['PUBLICATION_CITATION']=ref['citation']
        r['NUMERICAL_VERIFICATION_STATUS']='VERIFIED_FROM_PRIMARY_SOURCES'
        if r['NUMERICAL_STATUS'].startswith('COUNTS_CONCORDANT') and any(r[f'SOURCE_{i}_TYPE']=='ClinicalTrials.gov' for i in [1,2]):
            r['NUMERICAL_STATUS']='COUNTS_CONCORDANT; REGISTRY_PERCENT_DERIVED_NOT_SOURCE_DISPLAY'
        r['CONTEXT_ALIGNMENT_CERTIFIED']=False
        r['ORIGINAL_ROW_PRESERVED']=True
        r['INTERPRETATION_VERIFICATION_SCOPE']='Numerators and denominators checked; incomplete context is not certified as aligned.'
        if short in ['000003','000005','000040','000137']:
            r['VERIFICATION_STATUS']='VERIFIED'
            r['INTERPRETATION_VERIFICATION_SCOPE']='Documented population difference or explicit out-of-window exclusion verified; no numerical alignment claim.'
        for i in [1,2]:
            p=f'SOURCE_{i}_';typ=r[p+'TYPE']
            if typ=='ClinicalTrials.gov':
                assert r[p+'VERIFICATION_STATUS']=='NUMERATOR_DENOMINATOR_VERIFIED_FROM_ARCHIVED_JSON'
                r[p+'VERIFIED_LOCATOR']='resultsSection.adverseEventsModule.eventGroups; '+r[p+'TREATMENT_ARM_AND_DOSE'].split(';')[0]
                r[p+'VERIFIED_SOURCE_LABEL']='Participants affected by serious adverse events'
                r[p+'EVIDENCE_NOTE']='Archived JSON seriousNumAffected / seriousNumAtRisk; timeFrame and description read separately. Record update date is not the analysis cutoff.'
                r[p+'DOCUMENT_IDENTIFIER']=ref['nct_number']
                r[p+'CAUSALITY_ATTRIBUTION']='All reported serious events; not restricted to treatment-related events'
            else:
                stem=Path(r[p+'SOURCE_FILE']).stem;loc,evidence=source_index[stem]
                assert (HERE/r[p+'SOURCE_FILE']).exists()
                r[p+'VERIFIED_LOCATOR']=loc;r[p+'EVIDENCE_NOTE']=evidence
                r[p+'VERIFICATION_STATUS']='NUMERATOR_DENOMINATOR_AND_DISPLAY_CHECKED_FROM_PRIMARY_SOURCE'
                r[p+'VERIFIED_SOURCE_LABEL']=r[p+'TERM']
                r[p+'DOCUMENT_IDENTIFIER']=ref['doi'] if typ=='publication' else {'DREAMM2_FDA':'761158Orig1s000MultidisciplineR.pdf','EV201_FDA':'761137Orig1s000MultidisciplineR.pdf','IMMU13201_FDA':'761115Orig1s000MultidisciplineR.pdf','DESTINY_FDA':'761139Orig1s000MultidisciplineR.pdf'}[stem]
                r[p+'PERCENT_PROVENANCE']='Original display preserved and checked against source; calculation stored separately'
                if stem=='DREAMM2_FDA':r[p+'DATA_CUTOFF']='2019-06-21 (Table 27 source footnote)'
                if stem.startswith('innovaTV301'):
                    r[p+'DATA_CUTOFF']='2023-07-24 (publication)';r[p+'REPORTING_PERIOD']='Treatment-emergent through 30 days after last dose'
                if stem=='DESTINY_FDA':r[p+'ANALYSIS_POPULATION']='HER2-positive breast cancer 5.4 mg/kg safety pool; DS8201-A-J101 + DS8201-A-U201; N=234'
                if stem=='EV201_FDA':r[p+'REPORTING_PERIOD']='Treatment-emergent; Table 74 excludes the reported death >30 days after last dose'
        if short in ['000131','000143','000144']:
            r['SOURCE_1_VERIFIED_SOURCE_LABEL']={'000131':'Any treatment-related adverse event','000143':'Grade 3 or higher treatment-related adverse event','000144':'Grade 3 or higher adverse event'}[short]
            r['EXTRACTION_LABEL_CORRECTION']='Original term All adverse events retained. Verified table column/heading supplies attribution or severity; not a change to the original counts.'
        else:r['EXTRACTION_LABEL_CORRECTION']='None'
        if short=='000119':
            r['SOURCE_2_ADDITIONAL_REPORTED_PERCENT']=47
            r['REASON_FOR_DIFFERENCE']='Same count 47/99. FDA Table 27 displays 48%; FDA Table 30 and publication display 47%. Internal display inconsistency is documented; its cause is unresolved, and no rounding convention is inferred.'
        else:r['SOURCE_2_ADDITIONAL_REPORTED_PERCENT']=None
    for t in trials:
        reg=json.loads((ROOT/'data/raw/clinicaltrials'/(t['NCT_IDENTIFIER']+'.json')).read_text())
        status=reg['protocolSection']['statusModule'];date=status['lastUpdatePostDateStruct']['date']
        t['REGISTRY_LAST_UPDATE_POSTED']=date
        t['REGISTRY_FIRST_RESULTS_POSTED']=status.get('resultsFirstPostDateStruct',{}).get('date','UNKNOWN')
        t['CUTOFF_VERSION_STATUS']='POST_CUTOFF_VERSION; PRE_CUTOFF_FIELD_EQUIVALENCE_NOT_VERIFIED' if date>'2026-05-31' else 'ARCHIVED_VERSION_POSTED_BEFORE_TRIAL_IDENTIFICATION_CUTOFF'
        t['MAIN_LIMITATION']=re.sub(r'; archived registry version posted [^;]+ after identification cutoff; historical field equivalence NOT VERIFIED','',t['MAIN_LIMITATION'])
        evidence=date_evidence() if t['TRIAL_NAME']=='DREAMM-8' else None
        t['HISTORICAL_VERSION_POSTED']='2025-03-18' if evidence else None
        t['HISTORICAL_API_VERSION']=20 if evidence else None
        t['HISTORICAL_VERSION_EVIDENCE_FILE']='source_date_followup/DREAMM8_historical_comparison.json' if evidence else None
        t['HISTORICAL_AE_VERIFICATION']='Official historical and later modules identical; 2529 leaf fields. All six extracted n/N pairs match; three archived text-format differences documented.' if evidence else None
        if evidence:t['CUTOFF_VERSION_STATUS']='POST_CUTOFF_VERSION; PRE_CUTOFF_ADVERSE_EVENT_FIELD_EQUIVALENCE_VERIFIED'
        elif date>'2026-05-31':t['MAIN_LIMITATION']+='; archived registry version posted '+date+' after identification cutoff; historical field equivalence NOT VERIFIED'
        t['COVERAGE_INTERPRETATION']='Frozen extraction search-concept count; registry all-cause mortality was mapped historically to fatal-event search concept, not a verified fatal-adverse-event endpoint.'
        # No retained numerical publication-registry record is evidence of impossibility.
        t['PUBLICATION_REGISTRY_COMPARISON_POSSIBLE']=('NUMERIC RECORDS RETAINED; ALIGNMENT NOT FULLY VERIFIED' if t['TRIAL_NAME'] in ['DREAMM-2','innovaTV 301'] else 'NOT VERIFIED; BOTH EXTRACTIONS PRESENT' if t['REGISTRY_VALUES_EXTRACTED'] else 'NO REGISTRY CORE EXTRACTION IN ARCHIVE')
    save('record_audit.json',rows);save('trial_matrix.json',trials)
    sourcekeys=['RECORD_ID','TRIAL_NAME','NCT_IDENTIFIER','DOI','PUBLICATION_CITATION','SOURCE_PAIR','NUMERICAL_VERIFICATION_STATUS','COMPARABILITY_STATUS','VERIFICATION_STATUS','INTERPRETATION_VERIFICATION_SCOPE','REASON_FOR_DIFFERENCE','EXTRACTION_LABEL_CORRECTION']
    sourcekeys += [f'SOURCE_{i}_{k}' for i in [1,2] for k in ['TYPE','DOCUMENT_IDENTIFIER','SOURCE_FILE','ORIGINAL_LOCATOR','VERIFIED_LOCATOR','EVIDENCE_NOTE','VERIFICATION_STATUS','DATA_CUTOFF','REPORTING_PERIOD']]
    save('source_audit.json',[{k:r[k] for k in sourcekeys}|{'VERIFICATION_DATE':DATE} for r in rows])
    return rows,trials

def build_text(rows,trials):
    text=(PREV/'revised_manuscript.en.md').read_text()
    text=replace(text,'Supplementary Table S1 lists every trial and its archived source status.','Supplementary Table S1 lists every trial and its archived source status. The 31 May 2026 date defined trial identification; source versions were audited separately during revision on 25 October 2026. The archived DREAMM-8 registry version was posted on 12 June 2026, after that date, although results were first posted on 18 March 2025. Its field-level equivalence to a pre-cutoff version remains unverified.')
    text=replace(text,'The interpretation sequence began with a specific safety question.','A source inventory records which documents and extracted values are available. A numerical record pairs two extracted source rows; a source-context check examines whether those rows refer to the same group, outcome and observation period. Missing extraction means that no requested value was entered in the archive. A contextual comparison retains values with documented differences without interpreting their numerical difference as reporting error.\n\nThe interpretation sequence began with a specific safety question.')
    text=replace(text,'The previous source locators supported the publication and regulatory examples.','For all 28 records, numerators, denominators and displayed percentages were checked against the primary article or supplement, regulatory table or archived registry JSON. Numerical verification and reporting-context verification were recorded separately. Registry percentages calculated during extraction were distinguished from source-reported percentages. Incomplete context was not inferred from matching counts.')
    add17=('Within the 17 trials without structured regulatory extraction, publication and registry values were both available for 14 trials; LUMINOSITY, TROPION-Breast02 and CADENZA had publication extraction only. Nine trials were randomized controlled studies, including ASCENT, EMILIA and INO-VATE ALL. Supplementary Table S5 identifies these trials and their source contributions. In ASCENT and EMILIA, discontinuation, dose-reduction and grade 3 or higher outcomes were extracted from publications but not from the registry adverse-events extraction. Conversely, serious-event summaries were extracted from the registry but not the publication extraction for SG035-0003 and SG035-0004. These contrasts describe the archived extraction, without establishing absence from a complete source. Registry group descriptions and reporting periods added context, but did not establish aligned populations or cutoffs for all 14 trials. All 23 trials contribute to the source inventory; source availability, extraction completeness and suitability for numerical comparison are distinct properties.')
    assert 100<=words(add17)<=180
    anchor='Their regulatory extraction did not represent the safety reporting of the remaining trials.'
    text=replace(text,anchor,anchor+'\n\n'+add17)
    text=replace(text,'The regulatory review reported 233/234 (99.6%) in a broader safety pool.','The regulatory review reported 233/234 (99.6%) in a broader 5.4 mg/kg pool combining DS8201-A-J101 and DS8201-A-U201.')
    anchor='The metadata documented a broader reporting context; the archive did not determine how much of the count difference was attributable to follow-up duration.'
    text=replace(text,anchor,anchor+'\n\nEV-201 also had a non-identical serious-event count: the registry recorded 59/125 and regulatory Table 74 recorded 58/125. The registry reporting period extended to 62.36 months for cohort 1. This documented period difference did not establish the cause of the additional participant. In innovaTV 301, serious-event counts matched for tisotumab vedotin (82/250) and chemotherapy (94/239). Both sources specified collection through 30 days after the last dose, but the registry analysis cutoff could not be confirmed. These were numerical matches with incomplete reporting-context verification.')
    old='Their displayed percentages were 47% and 48%. In IMMU-132-01, discontinuation summaries had the same denominator of 108 but different numerators, three and five. The available archive did not resolve the cause of that difference.'
    new=('The publication displayed 47%; regulatory Table 27 displayed 48%, whereas regulatory Table 30 displayed 47%. The fraction 47/99 is 47.4747%, or 47.5% to one decimal place. The original displays were preserved; the internal regulatory inconsistency was not attributed to an undocumented rounding rule. In IMMU-132-01, the publication supplement recorded discontinuation in 3/108 participants and the regulatory review recorded 5/108. The review discussed individual discontinuation assessments but did not explicitly reconcile those five with the publication three. The cause remained unresolved.')
    text=replace(text,old,new)
    text=replace(text,'The paired regulatory table listed no drug-related treatment-emergent death.','The paired regulatory Table 74 listed no drug-related treatment-emergent death and explicitly excluded that death because it occurred more than 30 days after the last dose.')
    text=replace(text,'The supplement retains all 28 records with their original labels and revision annotations.','The supplement retains all 28 records with their original labels and separate numerical and context judgments. Eighteen had matching counts with incomplete context, eight had documented population or reporting-period differences, one had an unresolved count difference without an established explanation, and one was ineligible for aligned numerical interpretation. None was certified as fully context-aligned. These record counts describe the retained set and include repeated source contrasts.')
    text=replace(text,'Distinguishing counts from displayed percentages also gives readers a concrete check they can reproduce.','Distinguishing counts from displayed percentages also gives readers a concrete check they can reproduce. Matching counts alone do not establish common definitions, attribution or observation windows; incomplete metadata remain explicit in Supplementary Table S3.')
    text=replace(text,'The evaluation relied on aggregate reports and archived registry versions.','The evaluation relied on aggregate reports and archived registry versions. The post-cutoff DREAMM-8 version prevents a fully verified reconstruction of its registry fields as of 31 May 2026. The seven-concept coverage summaries also reproduce historical extraction labels: registry all-cause mortality had been mapped to the fatal-event search concept. Those counts describe extraction coverage, not harmonized fatal-event availability.')
    text=replace(text,'The accompanying editable supplement contains the 23-trial inventory, definitions of the original screening counts, the annotated 28-record ledger and worked-example source locators.','The accompanying editable supplement contains the 23-trial inventory, definitions of the original screening counts, the annotated 28-record ledger, worked-example source locators and trial-level source contributions and comparison barriers (Tables S1-S5).')
    text=replace(text,'Core outcomes refers to seven requested outcomes recorded in at least one source, without assuming identical definitions.','Core outcomes refers to seven historical extraction search concepts recorded in at least one source, without assuming identical definitions. Registry all-cause mortality was historically mapped to the fatal-event search concept; it is not a verified fatal-adverse-event endpoint. The frozen coverage counts reproduce that mapping.')
    text=replace(text,'Serious adverse events at 3.4 mg/kg: both sources 47/99; publication 47%, regulatory review 48%','Serious adverse events at 3.4 mg/kg: both sources 47/99; publication 47%, regulatory Table 27 48% and Table 30 47%')
    text=replace(text,'Counts agree; the displayed percentages differ. The calculated fraction is approximately 47.5%','Counts agree; an internal regulatory display inconsistency is documented. The calculated fraction is 47.4747%; no rounding explanation is inferred')
    text=replace(text,'| IMMU-132-01, unresolved count |','| EV-201, serious-event count | Registry 59/125; regulatory Table 74 58/125 | Counts differ; registry cohort 1 period extends to 62.36 months; cause of the extra participant unresolved |\n| innovaTV 301, numerical match | Publication and registry: tisotumab vedotin 82/250; chemotherapy 94/239 | Counts and arms match; a common registry analysis cutoff remains unverified |\n| IMMU-132-01, unresolved count |')
    text=replace(text,'Ann Oncol. 2026. doi: 10.1016/j.annonc.2026.03.008.','Ann Oncol. 2026;37(8):1066-1080. Published online 3 April 2026. doi: 10.1016/j.annonc.2026.03.008.')
    text=replace(text,'Chen and colleagues examined several safety outcomes across registry and publication locations, including withdrawals reported in participant flow.[2] Their approach shows why a search confined to one module can miss relevant information.','Chen and colleagues examined all-cause mortality, serious adverse events, adverse events and withdrawals due to adverse events across registry and publication records.[2] Their study illustrates the value of examining several safety outcomes.')
    text=text.replace('[6,11]','[6]')
    text=replace(text,'Regulatory guidance discusses evaluating individual studies and appropriately selected pools while accounting for differences in design and populations.[6]','Regulatory guidance discusses selected pools that account for design and population differences; FDA training also addresses safety-assessment design and analysis.[6,11]')
    if date_evidence():
        text=replace(text,'The archived DREAMM-8 registry version was posted on 12 June 2026, after that date, although results were first posted on 18 March 2025. Its field-level equivalence to a pre-cutoff version remains unverified.','The archived DREAMM-8 registry version was posted on 12 June 2026. We compared its adverse-event fields with the official historical version posted on 18 March 2025. The adverse-event modules in the two official versions were identical. All six extracted group-level numerators and denominators matched the historical version; three text-format differences between the archive and historical API representation involved only encoding or whitespace. Supplementary Table S5 documents this check.')
        text=replace(text,' The post-cutoff DREAMM-8 version prevents a fully verified reconstruction of its registry fields as of 31 May 2026.','')
    (HERE/'revised_manuscript.en.md').write_text(text)

    supp=(PREV/'supplementary_information.en.md').read_text()
    start=supp.index('## Supplementary Table S3');end=supp.index('## Registry metadata checks')
    lines=['## Supplementary Table S3 Record-level numerical and reporting-context audit','',
           'All 28 records and their original tiers are retained. The original term, percentage and locator remain unchanged in the audit workbook; verified labels and new judgments are additional fields. N = numerator/denominator agreement; C = count difference. A denotes matching counts with fully verified context; B matching counts with incompletely verified context; C an unresolved numerical difference; D a documented population or reporting-period difference; E ineligible for aligned numerical comparison. No record qualified for A. Counts do not certify common definitions, attribution or windows.','',
           'Registry percentages are calculated values from the extraction workflow, not percentages displayed in the archived JSON. They are marked calc below. A displayed percentage difference does not itself imply a numerator difference.','']
    for name in dict.fromkeys(r['TRIAL_NAME'] for r in rows):
        lines+=['### '+name,''];out=[]
        for r in rows:
            if r['TRIAL_NAME']!=name:continue
            values=[]
            for i in [1,2]:
                p=f'SOURCE_{i}_';v=r[p+'ORIGINAL_EXTRACTED_PERCENT']
                tag=' calc' if r[p+'TYPE']=='ClinicalTrials.gov' else ''
                values.append(f"{r[p+'NUMERATOR']}/{r[p+'DENOMINATOR']} ({v:g}%{tag})")
            outcome=r['OUTCOME_DEFINITION'].replace('_',' ')
            if r['RECORD_ID'][-6:] in ['000131','000143','000144']:outcome=r['SOURCE_1_VERIFIED_SOURCE_LABEL']
            out.append({'Record':r['RECORD_ID'][-6:],'Outcome':outcome,'Sources in order':r['SOURCE_PAIR'].replace('FDA review','regulatory'),'Source 1':values[0],'Source 2':values[1],'Original tier':'Primary' if r['ORIGINAL_TIER']=='primary_candidate' else 'Sensitivity','Judgment':r['REVISED_CATEGORY']+'; '+('N' if r['NUMERICAL_STATUS'].startswith('COUNTS_CONCORDANT') else 'C')})
        lines+=[table(out),'']
    lines+=['For 000119, 48% is the preserved FDA Table 27 display. FDA Table 30 displays 47% for the same 47/99. The calculated percentage is 47.4747%; the reason for the internal display inconsistency is unresolved. For 000131, 000143 and 000144, the original generic extracted term All adverse events is retained in the workbook; the verified table heading clarifies attribution or grade.','',
            'The eight D records comprise three population differences and five reporting-period contrasts. D does not establish the cause of a numerical difference. The unresolved count differences in 000106, 000114, 000120, 000123 and 000146 remain explicit; DREAMM-2 records repeat contrasts involving the same source groups. Record 000150 is C because its cause lacks an established reconciliation. Record 000137 is E and remains excluded from aligned numerical interpretation.','',
            'Numerators and denominators were source-checked for all 28 records. Four interpretation judgments (the three DESTINY-Breast01 population differences and the EV-201 out-of-period exclusion) are VERIFIED. The remaining 24 are PARTIALLY_VERIFIED for context, including all 18 matching-count records. These statuses do not claim independent duplicate adjudication.','',
            '## Supplementary Table S4 Primary-source checks for worked examples','']
    for ids,title in [(['000003','000005','000040'],'DESTINY-Breast01'),(['000119','000106','000114','000120','000123'],'DREAMM-2'),(['000137','000146'],'EV-201'),(['000150'],'IMMU-132-01'),(['000259','000262'],'innovaTV 301')]:
        rr=[r for r in rows if r['RECORD_ID'][-6:] in ids];r=rr[0]
        lines+=['### '+title+' ('+', '.join(ids)+')','',r['PUBLICATION_CITATION']+' Registry: '+r['NCT_IDENTIFIER']+'.','']
        seen=set()
        for x in rr:
            for i in [1,2]:
                p=f'SOURCE_{i}_';key=x[p+'SOURCE_FILE']
                if key in seen:continue
                seen.add(key)
                lines+=[x[p+'TYPE']+': '+x[p+'DOCUMENT_IDENTIFIER']+'. Locator: '+x[p+'VERIFIED_LOCATOR']+'. '+x[p+'EVIDENCE_NOTE'],'']
        for x in rr:
            lines+=[x['RECORD_ID']+': '+x['REASON_FOR_DIFFERENCE']+' Status: '+x['VERIFICATION_STATUS']+'.','']
    supp=supp[:start]+'\n'.join(lines)+'\n'+supp[end:]
    supp=replace(supp,'Core outcomes means the number of seven prespecified outcomes recorded in at least one source.','Core outcomes reproduces the historical seven-concept extraction count. Registry all-cause mortality was mapped historically to the fatal-event search concept; this does not establish a verified fatal-adverse-event endpoint.')
    s5=['## Supplementary Table S5 Source contributions and unresolved comparison barriers across all trials','',
        'P = extracted publication values; R = extracted registry values; F = extracted regulatory values. All 23 regulatory documents were retrieved. The presence of two extractions does not establish comparability. Outcome labels describe historical extraction search concepts; mortality means all-cause mortality wherever it came from the registry. All population, definition and reporting-period alignment fields remain NOT VERIFIED unless a specific worked-example difference is documented in S4.','']
    for fda in [True,False]:
        s5+=['### '+('Six trials with' if fda else 'Seventeen trials without')+' structured regulatory extraction','']
        entries=[]
        for t in trials:
            if t['FDA_CORE_VALUES_EXTRACTED']!=fda:continue
            trim=lambda s:s.replace('adverse_event_leading_to_discontinuation','discontinuation').replace('grade_3_or_higher_adverse_event','grade >=3').replace('serious_adverse_event','serious events').replace('fatal_adverse_event','fatal-event search label*').replace('any_adverse_event','any event').replace('_',' ')
            entries.append({'Trial':t['TRIAL_NAME'],'Extraction':('P; R' if t['REGISTRY_VALUES_EXTRACTED'] else 'P')+('; F' if fda else ''),'Design':'RCT' if t['DESIGN']=='randomized, controlled' else t['DESIGN'],'Publication-only extraction':trim(t['PUBLICATION_ONLY_EXTRACTED_CONCEPTS']),'Registry-only extraction':trim(t['REGISTRY_ONLY_EXTRACTED_CONCEPTS']),'Comparison barrier':'No registry core extraction' if not t['REGISTRY_VALUES_EXTRACTED'] else 'Aligned definitions, group and cutoff not fully verified'+('; post-cutoff registry version' if t['TRIAL_NAME']=='DREAMM-8' else '')})
        s5+=[table(entries),'']
    s5+=['*The registry fatal-event search label denotes historical mapping of all-cause mortality, not a verified fatal-adverse-event summary. RCT, randomized controlled trial. No new cross-source outcome pairs were generated. The electronic trial matrix records each registry group and reporting-period field.','',
          'DREAMM-8 (NCT04484623) first posted results on 18 March 2025. The archived record used in the extraction was last posted on 12 June 2026, after the 31 May 2026 identification cutoff. Its corresponding pre-cutoff adverse-event fields have not been independently recovered. Frozen counts are retained with this provenance limitation, not retrospectively changed.','']
    supp=supp.replace('## Reproducibility and provenance','\n'.join(s5)+'## Reproducibility and provenance')
    supp=supp.replace('The current revision scripts and outputs accompany this local supplement.','Verification date: 25 October 2026; this does not change cohort inclusion. The current revision scripts and outputs accompany this local supplement.')
    if date_evidence():
        supp=supp.replace('; post-cutoff registry version |','; historical adverse-event fields verified |')
        supp=replace(supp,'DREAMM-8 (NCT04484623) first posted results on 18 March 2025. The archived record used in the extraction was last posted on 12 June 2026, after the 31 May 2026 identification cutoff. Its corresponding pre-cutoff adverse-event fields have not been independently recovered. Frozen counts are retained with this provenance limitation, not retrospectively changed.','DREAMM-8 (NCT04484623): official historical version 21 (API studyVersion 20), submitted on 11 March 2025 and publicly posted on 18 March 2025, was compared with version 22 (API studyVersion 21), submitted on 25 May 2026 and posted on 12 June 2026. The adverse-event modules were identical across the two official versions (2529 leaf fields). The archived API snapshot differed from the historical representation in three text fields only: HTML entities, Markdown escapes and whitespace. All six extracted group-level n/N pairs matched the historical version: all-cause mortality 48/155 and 54/147; serious adverse events 95/150 and 65/145; other non-serious adverse events 148/150 and 130/145. The observation period and safety-population description were unchanged. This resolves availability of the extracted adverse-event fields before 31 May 2026; it does not establish alignment with the publication. No trial, record, selection rule or frozen count changed. Official history: https://clinicaltrials.gov/study/NCT04484623?tab=history&a=21. The historical responses, field comparison and hashes are preserved in the local source-date audit.')
    (HERE/'supplementary_information.en.md').write_text(supp)

    response=(PREV/'response.source.md').read_text()
    response=replace(response,'We provide editable clean and red-marked manuscripts, supplementary information and this combined cover letter and response, following the request to transcribe and answer the comments. Red text indicates new or substantially rewritten passages relative to the archived submission. Deletions are described in the responses and can be checked against the retained baseline.','We provide editable clean and color-marked manuscripts, supplementary information and this combined cover letter and response, following the request to transcribe and answer the comments. In the marked manuscript, red identifies retained changes from the first revision relative to the archived submission, and dark blue identifies substantive changes in this source-audit revision. Unchanged text remains black. These are text colors, not Word Track Changes; deletions are described in the responses and can be checked against the retained baseline.')
    response=response.replace('A separate red-marked Word manuscript','A separate color-marked Word manuscript')
    response=replace(response,'The remaining records are not presented as independently validated matches.','The remaining records are not presented as independently validated matches. In the final audit, all 28 numerators and denominators were checked against primary-source text, tables or archived JSON. Eighteen records had matching counts with incompletely verified context, eight had documented population or reporting-period differences, one had an unexplained numerical difference and one was ineligible for aligned numerical interpretation. No record was certified as fully context-aligned. S3 and the audit workbook distinguish numerical verification from context verification. The FDA Table 74 footnote directly supports the EV-201 exclusion; EV-201 59/125 versus 58/125 remains non-identical, and IMMU-132-01 3/108 versus 5/108 remains unresolved.')
    response=replace(response,'These are descriptions of extraction coverage, without significance tests or an inference that the six represent the other trials.','These are descriptions of extraction coverage, without significance tests or an inference that the six represent the other trials. Supplementary Table S5 and the electronic trial matrix now identify the 14 trials with both publication and registry extraction, the three without registry extraction, randomized designs and source-specific extraction contributions. The Results adds a short descriptive account. These additions do not create unsupported pairings. The medians reproduce historical search-concept mappings, including registry all-cause mortality under the fatal-event search label; they are not counts of harmonized fatal-adverse-event availability.')
    response=replace(response,'We also retain registry and publication contributions for trials without structured regulatory extraction.','We also retain registry and publication contributions for trials without structured regulatory extraction. Supplementary Table S5 documents those contributions and unresolved comparison barriers across every trial. The matching innovaTV 301 counts are explicitly classified as numerical matches with incomplete cutoff verification. The DREAMM-2 47/99 example now preserves both FDA displays: 48% in Table 27 and 47% in Table 30, without assuming a rounding rule.')
    response=replace(response,'The revised limitations explicitly restrict inference to this publication-anchored collection.','A separate date audit identified that the archived DREAMM-8 registry version was posted on 12 June 2026, although results were first posted in March 2025. Methods, limitations and Supplementary Table S5 disclose that the archived fields have not been verified against the version available by the 31 May 2026 identification cutoff. The frozen statistics were preserved, and no claim of a fully reconstructed cutoff-date registry dataset is made.\n\nThe revised limitations explicitly restrict inference to this publication-anchored collection.')
    if date_evidence():
        response=replace(response,'A separate date audit identified that the archived DREAMM-8 registry version was posted on 12 June 2026, although results were first posted in March 2025. Methods, limitations and Supplementary Table S5 disclose that the archived fields have not been verified against the version available by the 31 May 2026 identification cutoff. The frozen statistics were preserved, and no claim of a fully reconstructed cutoff-date registry dataset is made.','We also checked the source-version dates. For DREAMM-8, we retrieved official historical version 21 (API studyVersion 20), publicly posted on 18 March 2025, and compared it with version 22, posted on 12 June 2026. Their adverse-event modules were identical across 2529 leaf fields. All six group-level n/N pairs used in the extraction matched the historical version; three archive-versus-history text differences were formatting only. Methods and Supplementary Table S5 document that the extracted adverse-event fields were available before the 31 May 2026 trial-identification cutoff. No inclusion rule or frozen statistic changed. This version check does not establish cross-source reporting-context alignment.')
    (HERE/'response.source.md').write_text(response)
    save('bounded_changes.json',{'additional_17_trial_paragraph_words':words(add17),'introduction_changed':False,'new_analyses':False,'freeze_maintained':True,'new_source_issue':'RESOLVED: DREAMM-8 historical adverse-event fields verified' if date_evidence() else 'DREAMM-8 registry version posted after cutoff; field equivalence unverified'})
    return text,supp,response

def paragraphs(doc):
    out=list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:out.extend(c.paragraphs)
    return [p for p in out if p.text.strip()]
def tokenise(s):return re.findall(r'\s+|[A-Za-z0-9]+|[^\w\s]|_+',s)
def bestmatch(s,candidates):
    if s in candidates:return s
    # Paragraph correspondence is approximate; the color layer never changes text.
    ranked=sorted(candidates,key=lambda x:difflib.SequenceMatcher(None,s[:150],x[:150],autojunk=False).ratio(),reverse=True)[:5]
    return max(ranked,key=lambda x:difflib.SequenceMatcher(None,s,x,autojunk=False).ratio(),default='')
def changed_chars(new,old):
    nt=tokenise(new);ot=tokenise(old);changed=[True]*len(new);starts=[];pos=0
    for t in nt:starts.append(pos);pos+=len(t)
    for block in difflib.SequenceMatcher(None,ot,nt,autojunk=False).get_matching_blocks():
        for j in range(block.b,block.b+block.size):
            for k in range(starts[j],starts[j]+len(nt[j])):changed[k]=False
    return changed
def make_docs(text,supp,response):
    spec=importlib.util.spec_from_file_location('r1helpers',PREV/'revise_package.py');mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    targets={'CLEAN':text,'SUPPLEMENTARY':supp,'RESPONSE':response}
    for name,content in targets.items():mod.render_md(content,HERE/f'CT-26-0272_R2_FINAL_{name}.docx',manuscript=name=='CLEAN',supplement=name=='SUPPLEMENTARY')
    supplement=Document(HERE/'CT-26-0272_R2_FINAL_SUPPLEMENTARY.docx')
    from docx.shared import Pt
    for t in supplement.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_before=Pt(2)
                    p.paragraph_format.space_after=Pt(2)
    supplement.save(HERE/'CT-26-0272_R2_FINAL_SUPPLEMENTARY.docx')
    clean=Document(HERE/'CT-26-0272_R2_FINAL_CLEAN.docx')
    original=[p.text for p in paragraphs(Document(PREV/'baseline/submitted_manuscript.docx'))]
    previous=[p.text for p in paragraphs(Document(PREV/'CT-26-0272_revised_clean.docx'))]
    colorcounts=Counter()
    for p in paragraphs(clean):
        new=p.text;old=bestmatch(new,original);r1=bestmatch(new,previous)
        red=changed_chars(new,old);blue=changed_chars(new,r1)
        colors=['1F4E79' if b else 'C00000' if a else '000000' for a,b in zip(red,blue)]
        pos=0
        for run in list(p.runs):
            if not run.text:continue
            txt=run.text;chunks=[];start=0
            for k in range(1,len(txt)+1):
                if k==len(txt) or colors[pos+k]!=colors[pos+start]:chunks.append((txt[start:k],colors[pos+start]));start=k
            previous_r=run._r
            for k,(value,color) in enumerate(chunks):
                newr=copy.deepcopy(run._r);ts=newr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
                for t in ts:newr.remove(t)
                from docx.text.run import Run
                nr=Run(newr,p);nr.text=value;nr.font.color.rgb=RGBColor.from_string(color)
                previous_r.addnext(newr);previous_r=newr;colorcounts[color]+=len(value)
            pos+=len(txt);p._p.remove(run._r)
        assert p.text==new
    clean.save(HERE/'CT-26-0272_R2_FINAL_MARKED.docx')
    save('marking_checks.json',dict(colorcounts))
    return mod

def reviewer_matrix(response):
    letter=(PREV/'baseline/decision_letter.txt').read_text();norm=lambda s:re.sub(r'\s+',' ',s).strip()
    for start,end in [('Reviewer: 1','Reviewer: 2'),('Reviewer: 2','Associate Editor'),('Associate Editor','SUBMITTING A REVISED MANUSCRIPT')]:
        b=letter.rsplit(start,1)[1].split(end,1)[0]
        for line in b.splitlines():
            if line.strip():assert norm(line) in norm(response),line[:100]
    result=[];group=''
    for block in re.split(r'(?=^## |^### Comment )',response,flags=re.M):
        if block.startswith('## '):group=block.splitlines()[0][3:];continue
        if not block.startswith('### Comment '):continue
        header,body=block.split('\n',1);original,answer=body.split('Response.',1)
        exact=original.strip();status='FULLY_ADDRESSED';residual='No new validation or complete source capture is claimed.'
        if group=='Reviewer 1' and header.startswith('### Comment 1'):
            status='PARTIALLY_ADDRESSED';residual='All 23 trials retained; only six regulatory extractions and no independent duplicate adjudication. DREAMM-8 pre-cutoff adverse-event fields verified.' if date_evidence() else 'All 23 trials retained; only six regulatory extractions. DREAMM-8 pre-cutoff field equivalence and independent duplicate adjudication remain unavailable.'
        if group=='Reviewer 1' and header.startswith('### Comment 2'):
            status='PARTIALLY_ADDRESSED';residual='Units and records audited; full reporting-context alignment remains unverified for matching-count examples.'
        result.append({'Reviewer':group,'Comment number':header[12:],'Exact original comment':exact,'Core concern':header.split(' ',3)[-1],'Current response':answer.strip(),'Manuscript location':'Section and display locations stated in response; matches final headings','Supplement location':'Tables S1-S5; source workbooks' if 'Reviewer' in group or group=='Associate Editor' else 'Not applicable','Supporting evidence':'Original decision letter; final manuscript and supplement; analysis/record_audit.json; analysis/statistical_reconciliation.json','Resolution status':status,'Residual limitation':residual,'Original comment independently found':norm(exact) in norm(letter)})
    assert len(result)==16,len(result)
    save('reviewer_matrix.json',result);return result

def main():
    rows,trials=augment_audit();text,supp,response=build_text(rows,trials)
    make_docs(text,supp,response);matrix=reviewer_matrix(response)
    for quote in re.findall(r'^\*([^*]+)\*$',response,re.M):assert quote in text,quote
    abstract=text.split('\n## Abstract\n')[1].split('\n## Keywords\n')[0]
    intro=text.split('\n## Introduction\n')[1].split('\n## Methods\n')[0]
    discussion=text.split('\n## Discussion\n')[1].split('\n## Supplementary information\n')[0]
    body=text.split('\n## Introduction\n')[1].split('\n## Supplementary information\n')[0]
    metrics={'abstract_words_including_headings':words(abstract),'main_words_including_headings':words(body),'introduction_words':words(intro),'discussion_words':words(discussion),'discussion_to_introduction_ratio':round(words(discussion)/words(intro),2),'introduction_citation_occurrences':len(re.findall(r'\[[\d,]+\]',intro)),'discussion_citation_occurrences':len(re.findall(r'\[[\d,]+\]',discussion)),'references':38,'main_tables':3,'main_figures':2,'color_figures':2,'supplementary_tables':5,'reviewer_resolution':dict(Counter(r['Resolution status'] for r in matrix))}
    save('final_metrics.json',metrics);print(json.dumps(metrics,indent=2))

if __name__=='__main__':main()
